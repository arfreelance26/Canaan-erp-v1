"""Profitability → View → Excel: an easy-to-read audit report.

Sheets
  Summary        the big numbers + one trip worked out line by line
  Trips          one row per trip, read left to right like a bill:
                 Hire − Trip Expenses − Diesel − EMI − Tyres − Maintenance −
                 Documents = Profit after all costs
  Trip Expenses  every trip-sheet expense line, with a check that they add up
  Trucks         one row per truck for the period (actual bills / EMI / documents)
  Rates          each truck's per-km rates — every per-km cost above looks its
                 rate up here, so a wrong rate is fixed in one place

Money cells are live Excel formulas (click one to see how it's worked out).
openpyxl can't store a formula's result, so every value is also worked out
here: it feeds the in-app preview, and the numbers match the Profitability
page exactly (same Cost/KM, same per-km parts — see running_cost.py).
"""
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

IST = ZoneInfo("Asia/Kolkata")

# ── Look ─────────────────────────────────────────────────────────────────────
NAVY = "1B2B5E"
FILL = {
    "trip": "DCE6F2", "in": "E2EFDA", "out": "FCE4D6", "result": NAVY,
    "info": "EDEDED", "docs": "FFF2CC", "total": "F2F2F2", "review": "F2F2F2",
}
MONEY = '"₹"#,##0.00;[Red]-"₹"#,##0.00'
KM = "#,##0"
RATE = '"₹"0.00" /km"'
THIN = Side(style="thin", color="BFBFBF")


def _fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _show(v, fmt=None) -> str:
    """How a value looks in the preview (same formats as the file)."""
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        if fmt == MONEY:
            return f"{'-' if v < 0 else ''}₹{abs(v):,.2f}"
        if fmt == KM:
            return f"{v:,.0f}"
        if fmt == RATE:
            return f"₹{v:.2f} /km"
        return f"{v:,.2f}" if isinstance(v, float) else str(v)
    return str(v)


class Sheet:
    """Writes cells (value or formula) and keeps the matching preview grid."""

    def __init__(self, wb: Workbook, title: str):
        self.ws = wb.create_sheet(title)
        self.grid: dict[int, dict[int, str]] = {}
        # The worked-out number behind each cell (a formula cell holds only text).
        self.values: dict[tuple[int, int], float] = {}

    def put(self, r, c, value=None, formula=None, fmt=None, bold=False, fill=None, color=None, wrap=False, size=None):
        cell = self.ws.cell(row=r, column=c, value=formula if formula else value)
        if not formula and isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
            # Typed text (a customer name, a remark…) must never run as a formula.
            cell.data_type = "s"
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            self.values[(r, c)] = float(value)
        if fmt:
            cell.number_format = fmt
        if bold or color or size:
            cell.font = Font(bold=bold, color=color, size=size or 11)
        if fill:
            cell.fill = _fill(fill)
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        self.grid.setdefault(r, {})[c] = _show(value, fmt)
        return cell

    def value(self, r, c) -> float:
        return self.values.get((r, c), 0.0)

    def widths(self, widths):
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w

    def header(self, r, labels):
        for c, label in enumerate(labels, start=1):
            cell = self.put(r, c, label, bold=True, fill=NAVY, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[r].height = 30

    def groups(self, r, spans):
        """spans: [(first col, last col, label, fill key)] — the coloured band above the headers."""
        for c1, c2, label, key in spans:
            dark = key == "result"
            for c in range(c1, c2 + 1):
                self.ws.cell(row=r, column=c).fill = _fill(FILL[key])
            cell = self.ws.cell(row=r, column=c1, value=label)
            cell.font = Font(bold=True, color="FFFFFF" if dark else "1F2937")
            cell.alignment = Alignment(horizontal="center")
            if c2 > c1:
                self.ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)

    def preview(self, name, first_row=1):
        if not self.grid:
            return {"name": name, "rows": []}
        last_r = max(self.grid)
        last_c = max(max(cols) for cols in self.grid.values())
        rows = [[self.grid.get(r, {}).get(c, "") for c in range(1, last_c + 1)] for r in range(first_row, last_r + 1)]
        return {"name": name, "rows": rows}


EXPENSE_LINES = [
    ("driver_pay", "Driver Pay / Batta"), ("halt_pay", "Halt Pay"), ("toll_charges", "Toll"),
    ("port_pass_expense", "Port Pass"), ("weight_sheet_expense", "Weight Sheet"), ("mamol_expense", "Mamol"),
    ("claimable_mamol_expense", "Claimable Mamol"), ("traffic_rto_expense", "Traffic / RTO"),
    ("lift_on_off_expense", "Lift On/Off"), ("crane_operator_expense", "Crane Operator"),
    ("parking_expense", "Parking"), ("other_expenses", "Other"),
]
DOCS = [("rc", "RC"), ("fc", "FC"), ("road_tax", "Road Tax"), ("insurance", "Insurance"),
        ("national_permit", "National Permit"), ("local_permit", "Local Permit"),
        ("pollution_certificate", "Pollution Cert.")]

# Rates sheet columns (VLOOKUP index): A Truck ID … C Diesel, D EMI, E Tyres, F Maintenance, G Documents.
RATE_COL = {"diesel": 3, "emi": 4, "tyre": 5, "maintenance": 6, "compliance": 7}


def _lookup(km_ref, truck_ref, part):
    return f"{km_ref}*VLOOKUP({truck_ref},Rates!$A:$G,{RATE_COL[part]},FALSE)"


def _d(s: str | None) -> str:
    return f"{s[8:10]}-{s[5:7]}-{s[:4]}" if s else ""


def build(db: Session, start: date, end: date, mode: str):
    """Returns (xlsx bytes, filename, preview sheets)."""
    import pl_compute

    # Every number comes from pl_compute — the same figures the page shows.
    # Full precision, so the preview equals what Excel's own formulas give.
    R = pl_compute.rates(db, mode)
    res = pl_compute.compute(db, start, end, mode, R=R, rounded=False)
    trucks, trips = res["trucks"], res["trips"]
    cpk, diesel_pk, parts = R["cpk"], R["diesel"], R["split"]

    wb = Workbook()
    wb.remove(wb.active)
    summary = Sheet(wb, "Summary")
    trip_s = Sheet(wb, "Trips")
    exp_s = Sheet(wb, "Trip Expenses")
    truck_s = Sheet(wb, "Trucks")
    rate_s = Sheet(wb, "Rates")

    # ── Rates ────────────────────────────────────────────────────────────────
    rate_s.header(1, ["Truck ID", "Registration", "Diesel", "EMI", "Tyres", "Maintenance", "Documents",
                      "Total per km", "Cost/KM used", "Check", "Where to fix a wrong rate"])
    rates: dict[str, dict] = {}
    where = f"Running Cost Calculator → {mode}"
    for i, e in enumerate(trucks):
        r = i + 2
        tid = e["truck_id"]
        p = parts.get(tid) or {}
        known = bool(p.get("parts_known"))
        rate = {
            "diesel": diesel_pk.get(tid),
            "emi": p.get("emi") if known else None, "tyre": p.get("tyre") if known else None,
            "maintenance": p.get("maintenance") if known else None, "compliance": p.get("compliance") if known else None,
        }
        rates[tid] = {**rate, "known": known, "cpk": cpk.get(tid)}
        total = sum(v or 0 for v in rate.values())
        used = cpk.get(tid)
        rate_s.put(r, 1, tid, bold=True)
        rate_s.put(r, 2, e.get("registration_number") or "")
        for c, k in enumerate(("diesel", "emi", "tyre", "maintenance", "compliance"), start=3):
            rate_s.put(r, c, rate[k], fmt=RATE)
        rate_s.put(r, 8, total, formula=f"=SUM(C{r}:G{r})", fmt=RATE)
        rate_s.put(r, 9, used, fmt=RATE)
        if used is None:
            rate_s.put(r, 10, "Not set", color="C00000", bold=True)
            rate_s.put(r, 11, f"Set it in {where}")
        else:
            ok = known and abs(total - used) < 0.0005
            rate_s.put(r, 10, "OK" if ok else "CHECK", formula=f'=IF(ABS(H{r}-I{r})<0.0005,"OK","CHECK")', bold=True)
            rate_s.put(r, 11, where if known else f"{where} — save the Manual tab to split the rates")
    rate_s.widths([12, 15, 13, 13, 13, 14, 13, 14, 14, 9, 48])
    rate_s.ws.freeze_panes = "C2"
    n_r = len(trucks) + 1
    rate_s.ws.conditional_formatting.add(f"J2:J{n_r}", CellIsRule(operator="equal", formula=['"OK"'], font=Font(color="00873C", bold=True)))
    rate_s.ws.conditional_formatting.add(f"J2:J{n_r}", CellIsRule(operator="equal", formula=['"CHECK"'], font=Font(color="C00000", bold=True)))

    # ── Trips ────────────────────────────────────────────────────────────────
    trip_s.groups(1, [(1, 7, "TRIP", "trip"), (8, 8, "MONEY IN", "in"), (9, 14, "MONEY OUT", "out"),
                      (15, 15, "RESULT", "result"), (16, 17, "INFO", "info")])
    trip_s.header(2, ["Trip ID", "Date", "Truck ID", "Registration", "Customer", "Route", "KM",
                      "Hire", "Trip Expenses", "Diesel", "EMI", "Tyres", "Maintenance", "Documents",
                      "Profit after all costs", "Status", "Notes"])
    first, last = 4, 3 + len(trips)
    example = None
    for i, t in enumerate(trips):
        r = first + i
        tid = t["truck_id"]
        rt = rates.get(tid) or {}
        km = float(t.get("total_km") or 0)
        hire, texp = float(t["hire_amount"]), float(t["total_expense"])
        counted = not t.get("excluded_from_totals")
        notes = []
        if not counted:
            notes.append(f"Under audit review{': ' + t['flagged_remark'] if t.get('flagged_remark') else ''}")
        grey = None if counted else "808080"
        trip_s.put(r, 1, t.get("trip_code") or "—", bold=True, color=grey)
        trip_s.put(r, 2, _d(t.get("trip_sheet_date")), color=grey)
        trip_s.put(r, 3, tid, color=grey)
        trip_s.put(r, 4, t["registration_number"], color=grey)
        trip_s.put(r, 5, t.get("customer_name") or "—", color=grey)
        route = f"{t.get('from_location')} → {t.get('to_location')}" if t.get("from_location") and t.get("to_location") else "—"
        trip_s.put(r, 6, route, color=grey)
        trip_s.put(r, 7, km, fmt=KM, color=grey)
        trip_s.put(r, 8, hire, fmt=MONEY, color=grey)
        trip_s.put(r, 9, texp, fmt=MONEY, color=grey)

        m = t["money"]
        c = rt.get("cpk")
        d_rate = rt.get("diesel")
        diesel = m["diesel"] or 0.0
        trip_s.put(r, 10, diesel, formula="=" + _lookup(f"G{r}", f"C{r}", "diesel"), fmt=MONEY, color=grey)
        if d_rate is None:
            notes.append("No diesel rate")
        profit = None
        if c is None or km <= 0:
            if c is None:
                notes.append(f"No Cost/KM in {mode} mode — profit can't be worked out")
        elif m["split"]:
            for col, part, label in ((11, "emi", "EMI"), (12, "tyre", "tyre"), (13, "maintenance", "maintenance"), (14, "compliance", "documents")):
                trip_s.put(r, col, m["split"][part] or 0.0, formula="=" + _lookup(f"G{r}", f"C{r}", part), fmt=MONEY, color=grey)
                if rt.get(part) is None:
                    notes.append(f"No {label} rate")
            profit = m["final_profit"]
            trip_s.put(r, 15, profit, formula=f"=H{r}-I{r}-J{r}-K{r}-L{r}-M{r}-N{r}", fmt=MONEY, bold=True, color=grey)
        else:
            lump = m["truck_expenses"]
            profit = m["final_profit"]
            trip_s.put(r, 15, profit, formula=f"=H{r}-I{r}-J{r}-{lump:.6f}", fmt=MONEY, bold=True, color=grey)
            notes.append(f"Running costs ₹{lump:,.2f} not itemised — save the Manual tab in the Running Cost Calculator")
        trip_s.put(r, 16, "Counted" if counted else "Under review", color=grey)
        trip_s.put(r, 17, "; ".join(notes), color=grey)
        if not counted:
            for col in range(1, 18):
                trip_s.ws.cell(row=r, column=col).fill = _fill(FILL["review"])
        # The worked example: a counted trip with the most rates set (then the
        # most profitable) — a rate that isn't set shows as "counted as 0".
        if counted and profit is not None and rt.get("known") and d_rate is not None:
            score = (sum(rt.get(k) is not None for k in ("emi", "tyre", "maintenance", "compliance")), km >= 100, profit)
            if example is None or score > example[0]:
                example = (score, r, profit, t, km, hire, texp, diesel, rt)

    # Totals (row 3) — counted trips only, as on the page.
    trip_s.put(3, 1, "TOTAL (counted trips)", bold=True, fill=FILL["total"])
    for col in range(2, 18):
        trip_s.ws.cell(row=3, column=col).fill = _fill(FILL["total"])
    counted_rows = [first + i for i, t in enumerate(trips) if not t.get("excluded_from_totals")]
    for col, fmt in ((7, KM), (8, MONEY), (9, MONEY), (10, MONEY), (11, MONEY), (12, MONEY), (13, MONEY), (14, MONEY), (15, MONEY)):
        L = get_column_letter(col)
        total = sum(trip_s.value(rr, col) for rr in counted_rows)
        trip_s.put(3, col, total, formula=f'=SUMIF($P${first}:$P${last},"Counted",{L}{first}:{L}{last})' if trips else None,
                   fmt=fmt, bold=True, fill=FILL["total"])
    trip_s.widths([11, 11, 11, 14, 24, 32, 8, 13, 13, 12, 12, 12, 13, 12, 17, 13, 55])
    trip_s.ws.freeze_panes = "B4"
    if trips:
        trip_s.ws.auto_filter.ref = f"A2:Q{last}"
        rng = f"O{first}:O{last}"
        trip_s.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER(O{first}),O{first}>=0)'], fill=_fill("E2EFDA"), font=Font(color="00612E", bold=True)))
        trip_s.ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=_fill("FDE2E1"), font=Font(color="C00000", bold=True)))

    # ── Trip Expenses ────────────────────────────────────────────────────────
    labels = [l for _, l in EXPENSE_LINES]
    exp_s.header(1, ["Trip ID", "Date", "Truck ID", *labels, "Lines Total", "Trip Expenses (on Trips sheet)", "Check"])
    n_lines = len(EXPENSE_LINES)
    c_total, c_rec, c_chk = 4 + n_lines, 5 + n_lines, 6 + n_lines
    L1, Ln = get_column_letter(4), get_column_letter(3 + n_lines)
    Lt, Lr = get_column_letter(c_total), get_column_letter(c_rec)
    for i, t in enumerate(trips):
        r = 3 + i
        exp_s.put(r, 1, t.get("trip_code") or "—", bold=True)
        exp_s.put(r, 2, _d(t.get("trip_sheet_date")))
        exp_s.put(r, 3, t["truck_id"])
        lines = [float((t.get("expense_lines") or {}).get(k) or 0) for k, _ in EXPENSE_LINES]
        for j, v in enumerate(lines):
            exp_s.put(r, 4 + j, v or None, fmt=MONEY)
        total = sum(lines)
        exp_s.put(r, c_total, total, formula=f"=SUM({L1}{r}:{Ln}{r})", fmt=MONEY, bold=True)
        exp_s.put(r, c_rec, float(t["total_expense"]), fmt=MONEY)
        ok = abs(total - float(t["total_expense"])) < 0.01
        exp_s.put(r, c_chk, "OK" if ok else "MISMATCH", formula=f'=IF(ABS({Lt}{r}-{Lr}{r})<0.01,"OK","MISMATCH")', bold=True)
    exp_s.put(2, 1, "TOTAL (all trips)", bold=True, fill=FILL["total"])
    for col in range(2, c_chk + 1):
        exp_s.ws.cell(row=2, column=col).fill = _fill(FILL["total"])
    e_last = 2 + len(trips)
    for col in range(4, c_rec + 1):
        L = get_column_letter(col)
        exp_s.put(2, col, sum(exp_s.value(rr, col) for rr in range(3, e_last + 1)),
                  formula=f"=SUM({L}3:{L}{e_last})" if trips else None, fmt=MONEY, bold=True, fill=FILL["total"])
    exp_s.widths([11, 11, 11, *[13] * n_lines, 14, 17, 11])
    exp_s.ws.freeze_panes = "B3"
    if trips:
        chk = f"{get_column_letter(c_chk)}3:{get_column_letter(c_chk)}{e_last}"
        exp_s.ws.conditional_formatting.add(chk, CellIsRule(operator="equal", formula=['"OK"'], font=Font(color="00873C", bold=True)))
        exp_s.ws.conditional_formatting.add(chk, CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=_fill("FDE2E1"), font=Font(color="C00000", bold=True)))

    # ── Trucks ───────────────────────────────────────────────────────────────
    n_docs = len(DOCS)
    truck_s.groups(1, [(1, 5, "TRUCK", "trip"), (6, 6, "MONEY IN", "in"), (7, 12, "MONEY OUT", "out"),
                       (13, 13, "RESULT", "result"), (14, 13 + n_docs, "DOCUMENTS — BREAKDOWN", "docs"),
                       (14 + n_docs, 14 + n_docs, "INFO", "info")])
    truck_s.header(2, ["Truck ID", "Registration", "Status", "Trips", "KM", "Hire", "Trip Expenses",
                       "Diesel", "Tyres", "Maintenance (bills)", "EMI (share)", "Documents",
                       "Profit after all costs", *[l for _, l in DOCS], "Notes"])
    D1, Dn = get_column_letter(14), get_column_letter(13 + n_docs)
    t_first, t_last = 4, 3 + len(trucks)
    for i, e in enumerate(trucks):
        r = t_first + i
        tid = e["truck_id"]
        rt = rates.get(tid) or {}
        km = float(e.get("total_km") or 0)
        hire, texp = float(e["total_hire_amount"]), float(e["trip_expenses"])
        maint, emi = float(e["maintenance_expenses"]), float(e["emi_share"])
        docs = {k: float((e.get("document_breakdown") or {}).get(k) or 0) for k, _ in DOCS}
        m = e["money"]
        diesel, tyres = m["diesel"] or 0.0, m["tyres"] or 0.0
        doc_total = sum(docs.values())
        notes = []
        if rt.get("diesel") is None:
            notes.append("No diesel rate (counted as 0)")
        if rt.get("tyre") is None:
            notes.append("No tyre rate (counted as 0)")
        if int(e.get("flagged_trip_count") or 0):
            notes.append(f"{e['flagged_trip_count']} trip(s) under audit review — not counted")
        truck_s.put(r, 1, tid, bold=True)
        truck_s.put(r, 2, e.get("registration_number") or "")
        truck_s.put(r, 3, e.get("disposal") or "Active")
        truck_s.put(r, 4, int(e.get("trip_count") or 0))
        truck_s.put(r, 5, km, fmt=KM)
        truck_s.put(r, 6, hire, fmt=MONEY)
        truck_s.put(r, 7, texp, fmt=MONEY)
        truck_s.put(r, 8, diesel, formula="=" + _lookup(f"E{r}", f"A{r}", "diesel"), fmt=MONEY)
        truck_s.put(r, 9, tyres, formula="=" + _lookup(f"E{r}", f"A{r}", "tyre"), fmt=MONEY)
        truck_s.put(r, 10, maint, fmt=MONEY)
        truck_s.put(r, 11, emi, fmt=MONEY)
        truck_s.put(r, 12, doc_total, formula=f"=SUM({D1}{r}:{Dn}{r})", fmt=MONEY)
        truck_s.put(r, 13, m["net"],
                    formula=f"=F{r}-G{r}-H{r}-I{r}-J{r}-K{r}-L{r}", fmt=MONEY, bold=True)
        for j, (k, _) in enumerate(DOCS):
            truck_s.put(r, 14 + j, docs[k] or None, fmt=MONEY)
        truck_s.put(r, 14 + n_docs, "; ".join(notes))
    truck_s.put(3, 1, "TOTAL", bold=True, fill=FILL["total"])
    for col in range(2, 15 + n_docs):
        truck_s.ws.cell(row=3, column=col).fill = _fill(FILL["total"])
    for col in [4, 5, *range(6, 14 + n_docs)]:
        L = get_column_letter(col)
        fmt = None if col == 4 else KM if col == 5 else MONEY
        truck_s.put(3, col, sum(truck_s.value(rr, col) for rr in range(t_first, t_last + 1)),
                    formula=f"=SUM({L}{t_first}:{L}{t_last})" if trucks else None, fmt=fmt, bold=True, fill=FILL["total"])
    truck_s.widths([11, 15, 9, 7, 9, 14, 14, 13, 12, 14, 12, 12, 17, *[12] * n_docs, 45])
    truck_s.ws.freeze_panes = "C4"
    if trucks:
        truck_s.ws.auto_filter.ref = f"A2:{get_column_letter(14 + n_docs)}{t_last}"
        rng = f"M{t_first}:M{t_last}"
        truck_s.ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=_fill("E2EFDA"), font=Font(color="00612E", bold=True)))
        truck_s.ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=_fill("FDE2E1"), font=Font(color="C00000", bold=True)))

    # ── Summary ──────────────────────────────────────────────────────────────
    s = summary
    s.put(1, 1, "Profitability Report", bold=True, size=16, color=NAVY)
    s.put(2, 1, f"{_d(start.isoformat())} to {_d(end.isoformat())}  ·  Cost/KM mode: {mode}")
    s.put(3, 1, f"Downloaded {datetime.now(IST).strftime('%d-%m-%Y %I:%M %p')}", color="808080")

    def tot(col):
        return trip_s.value(3, col)

    n_counted = len(counted_rows)
    n_review = len(trips) - n_counted
    losses = sum(1 for rr in counted_rows if (rr, 15) in trip_s.values and trip_s.value(rr, 15) < 0)
    running = tot(11) + tot(12) + tot(13) + tot(14)
    s.put(5, 1, "THE BIG PICTURE (counted trips)", bold=True, color="FFFFFF", fill=NAVY)
    s.put(5, 2, "", fill=NAVY)
    big = [
        ("Trips", n_counted, f'=COUNTIF(Trips!$P${first}:$P${last},"Counted")', None),
        ("Trips under audit review (not counted)", n_review, f'=COUNTIF(Trips!$P${first}:$P${last},"Under review")', None),
        ("Hire (money in)", tot(8), "=Trips!H3", MONEY),
        ("Trip expenses", tot(9), "=Trips!I3", MONEY),
        ("Diesel", tot(10), "=Trips!J3", MONEY),
        ("Running costs (EMI + Tyres + Maintenance + Documents)", running, "=SUM(Trips!K3:N3)", MONEY),
        ("Profit after all costs", tot(15), "=Trips!O3", MONEY),
        ("Trips that made a loss", losses, f'=COUNTIFS(Trips!$O${first}:$O${last},"<0",Trips!$P${first}:$P${last},"Counted")', None),
    ]
    for j, (label, v, f, fmt) in enumerate(big):
        r = 6 + j
        is_profit = label.startswith("Profit")
        s.put(r, 1, label, bold=is_profit)
        s.put(r, 2, v, formula=f if trips else None, fmt=fmt, bold=True,
              color=("00612E" if (v or 0) >= 0 else "C00000") if is_profit else None)

    r = 15
    s.put(r, 1, "HOW ONE TRIP IS WORKED OUT", bold=True, color="FFFFFF", fill=NAVY)
    s.put(r, 2, "", fill=NAVY)
    if example:
        _, er, profit, t, km, hire, texp, diesel, rt = example
        route = f"{t.get('from_location')} → {t.get('to_location')}"

        def per_km(part, label, col):
            rate = rt.get(part)
            how = f"{km:,.0f} km × ₹{rate:.2f} per km" if rate is not None else "no rate set — counted as 0"
            return (f"− {label}  ({how})", km * (rate or 0), f"=Trips!{col}{er}", MONEY)
        steps = [
            ("Trip", f"{t.get('trip_code')}  ·  {route}  ·  {km:,.0f} km", None, None),
            ("Hire (money in)", hire, f"=Trips!H{er}", MONEY),
            ("− Trip expenses (driver pay, toll, port pass, …)", texp, f"=Trips!I{er}", MONEY),
            (f"− Diesel  ({km:,.0f} km × ₹{rt['diesel']:.2f} per km)", diesel, f"=Trips!J{er}", MONEY),
            per_km("emi", "EMI", "K"),
            per_km("tyre", "Tyres", "L"),
            per_km("maintenance", "Maintenance", "M"),
            per_km("compliance", "Documents", "N"),
            ("Profit after all costs", profit, f"=Trips!O{er}", MONEY),
        ]
        for j, (label, v, f, fmt) in enumerate(steps):
            rr = r + 1 + j
            last_step = j == len(steps) - 1   # (a label starting "=" would become a formula)
            s.put(rr, 1, label, bold=last_step or j == 0)
            s.put(rr, 2, v, formula=f, fmt=fmt, bold=last_step)
            if last_step:
                s.ws.cell(row=rr, column=1).border = Border(top=THIN)
                s.ws.cell(row=rr, column=2).border = Border(top=THIN)
        r = r + 1 + len(steps)
    else:
        s.put(r + 1, 1, "No counted trip in this period has its truck's rates set, so there's no example to show.")
        r += 2

    r += 1
    s.put(r, 1, "WHAT'S IN THIS FILE", bold=True, color="FFFFFF", fill=NAVY)
    s.put(r, 2, "", fill=NAVY)
    guide = [
        ("Trips", "One row per trip. Read it left to right: Hire − every cost = Profit."),
        ("Trip Expenses", "Every expense line of each trip sheet. Check says OK when they add up."),
        ("Trucks", "One row per truck for the whole period, using the actual bills, EMI and documents."),
        ("Rates", "Each truck's cost per km. Every per-km cost looks its rate up here — fix a wrong rate once."),
        ("Tip", "Click any amount to see how it's worked out (it's a formula)."),
        ("Good to know", "Trips charge maintenance, EMI and documents at an average per km; Trucks use the real "
                         "bills in this period — so the two sheets' totals differ, and both are right."),
        ("Grey rows", "Trips under audit review — listed, but not counted until the review is done."),
    ]
    for j, (a, b) in enumerate(guide):
        s.put(r + 1 + j, 1, a, bold=True)
        s.put(r + 1 + j, 2, b, wrap=True)
    s.widths([58, 90])

    for sh in (trip_s, exp_s, truck_s, rate_s):
        sh.ws.sheet_view.zoomScale = 90
    wb.calculation.fullCalcOnLoad = True

    buf = BytesIO()
    wb.save(buf)
    filename = f"Profitability_{_d(start.isoformat())}_to_{_d(end.isoformat())}_{mode}.xlsx"
    preview = [
        summary.preview("Summary"),
        trip_s.preview("Trips", first_row=2),
        exp_s.preview("Trip Expenses"),
        truck_s.preview("Trucks", first_row=2),
        rate_s.preview("Rates"),
    ]
    return buf.getvalue(), filename, preview
