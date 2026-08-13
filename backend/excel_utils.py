import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from fastapi.responses import StreamingResponse
from sqlalchemy import LargeBinary
from sqlalchemy.orm import Session


def _serialize(value: Any):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value
    return value


def exportable_columns(model) -> list[str]:
    """All column names on a model except BLOB/binary columns, which aren't spreadsheet-friendly."""
    return [c.name for c in model.__table__.columns if not isinstance(c.type, LargeBinary)]


def header_for(column_name: str) -> str:
    return column_name.replace("_", " ").title()


def rows_from_model(db: Session, model, order_by=None) -> tuple[list[str], list[list[Any]]]:
    """Query every row of `model` and return (headers, rows) using every schema column (minus BLOBs)."""
    columns = exportable_columns(model)
    headers = [header_for(c) for c in columns]
    q = db.query(model)
    if order_by is not None:
        q = q.order_by(order_by)
    rows = [[_serialize(getattr(obj, col)) for col in columns] for obj in q.all()]
    return headers, rows


def build_excel_response(sheets: list[tuple[str, list[str], list[list[Any]]]], filename: str) -> StreamingResponse:
    """Build a multi-sheet .xlsx workbook and return it as a downloadable FastAPI response."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        safe_name = "".join(ch for ch in name if ch not in ':\\/?*[]')[:31] or "Sheet1"
        ws = wb.create_sheet(title=safe_name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
            # Force Indian DD-MM-YYYY display on any date/datetime cells.
            excel_row = ws.max_row
            for col_idx, value in enumerate(row, start=1):
                if isinstance(value, datetime):
                    ws.cell(row=excel_row, column=col_idx).number_format = "DD-MM-YYYY HH:MM:SS"
                elif isinstance(value, date):
                    ws.cell(row=excel_row, column=col_idx).number_format = "DD-MM-YYYY"
        for i, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(len(str(header)) + 2, 12)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_filename = filename if filename.endswith(".xlsx") else f"{filename}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_filename}"'},
    )
