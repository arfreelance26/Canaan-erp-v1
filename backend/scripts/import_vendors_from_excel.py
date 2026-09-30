"""One-time script: imports vendor records from an Excel sheet into the
`vendors` table (the "Our Vendors" page).

Expected sheet layout (matches "Vendor list given by Jebarson.xlsx"):
    Row 2 header:  | (blank) | Name of the vendor | Vendor Category | Contact number | GSTIN | PAN | Email | ADDRESS |
    Data rows start at row 4. Column A (a per-category running number) is ignored.

Behaviour:
- Every data row becomes its own Vendor record, even if the same vendor name
  appears more than once under a different category (matches how the source
  sheet lists a vendor once per service they provide, e.g. "JAYARAJ TRUCKS &
  BUSES PVT. LTD." appears under Adblue vendor, Spares, and Mechanical &
  Electrical Workshop as three separate rows).
- A row with no category filled in is imported with an empty category.
- Vendor Category text is matched case-insensitively against whatever is
  already in the `vendor_categories` table; a small alias map fixes known
  typos in the sheet (e.g. "Worshop" -> "Workshop"). If, after that, the
  resolved category still doesn't exist, this script creates it (same as
  add_vendor_categories.py) so the import never fails for a missing category.
- Idempotent: before inserting, it checks whether a non-deleted vendor with
  the same (name, resolved category) already exists and skips it if so —
  safe to re-run without doubling up the list.

Usage (from the backend/ directory, with the same environment/DB config the
target database uses):
    python scripts/import_vendors_from_excel.py "<path to xlsx>" --dry-run
    python scripts/import_vendors_from_excel.py "<path to xlsx>"
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from database import SessionLocal
import models

DATA_START_ROW = 4
COL_NAME = 1  # index into the row tuple (A, B, C, D, E, F, G, H)
COL_CATEGORY = 2
COL_CONTACT = 3
COL_GSTIN = 4
COL_PAN = 5
COL_EMAIL = 6
COL_ADDRESS = 7

# Known typos/wording differences between the sheet and the canonical
# category names — matched after lowercasing, so case doesn't matter here.
CATEGORY_ALIASES = {
    "mechanical & electrical worshop": "MECHANICAL & ELECTRICAL WORKSHOP",
}


def clean_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


CONTACT_MAX_LEN = 20  # matches Vendor.contact_number's VARCHAR(20) column


def clean_contact(value) -> tuple[str, str]:
    """Returns (primary_number, note). Some sheet rows list two or three
    numbers separated by commas/spaces, which overflow the single-number
    contact_number column (VARCHAR(20)) — only the first is kept, and `note`
    is non-empty when something had to be dropped so the caller can warn."""
    if value is None:
        return "", ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    raw = str(value).strip()
    if not raw:
        return "", ""
    tokens = [t for t in re.split(r"[,\s]+", raw) if t]
    primary = tokens[0][:CONTACT_MAX_LEN]
    note = raw if len(tokens) > 1 or len(tokens[0]) > CONTACT_MAX_LEN else ""
    return primary, note


def resolve_category(raw, existing_by_lower: dict, db, created: list) -> str:
    """Returns the exact category name to store on the vendor, creating a new
    VendorCategory row (and recording it in `created`) if the resolved name
    doesn't exist under any known spelling yet. Returns "" for a blank input."""
    raw = clean_str(raw)
    if not raw:
        return ""
    key = raw.lower()
    if key in existing_by_lower:
        return existing_by_lower[key]
    canonical = CATEGORY_ALIASES.get(key, raw)
    canonical_key = canonical.lower()
    if canonical_key in existing_by_lower:
        existing_by_lower[key] = existing_by_lower[canonical_key]
        return existing_by_lower[canonical_key]
    # Not found under any known spelling — create it.
    db.add(models.VendorCategory(name=canonical))
    db.flush()
    existing_by_lower[canonical_key] = canonical
    existing_by_lower[key] = canonical
    created.append(canonical)
    return canonical


def main(xlsx_path: str, dry_run: bool):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    db = SessionLocal()
    added = []
    skipped = []
    categories_created = []
    contact_warnings = []
    try:
        existing_categories_by_lower = {
            c.name.strip().lower(): c.name
            for c in db.query(models.VendorCategory).all()
        }
        existing_vendor_keys = {
            (v.name.strip().lower(), (v.category or "").strip().lower())
            for v in db.query(models.Vendor).filter(models.Vendor.deleted_at.is_(None)).all()
        }

        for row in ws.iter_rows(min_row=DATA_START_ROW, max_row=ws.max_row, values_only=True):
            name = clean_str(row[COL_NAME])
            if not name:
                continue

            category = resolve_category(row[COL_CATEGORY], existing_categories_by_lower, db, categories_created)
            key = (name.lower(), category.lower())
            if key in existing_vendor_keys:
                skipped.append((name, category))
                continue

            contact, contact_note = clean_contact(row[COL_CONTACT])
            if contact_note:
                contact_warnings.append((name, contact_note, contact))

            if not dry_run:
                db.add(models.Vendor(
                    name=name,
                    category=category or None,
                    contact_number=contact or None,
                    gstin=clean_str(row[COL_GSTIN]).upper() or None,
                    pan=clean_str(row[COL_PAN]).upper() or None,
                    email=clean_str(row[COL_EMAIL]) or None,
                    address=clean_str(row[COL_ADDRESS]) or None,
                ))
            existing_vendor_keys.add(key)
            added.append((name, category))

        if dry_run:
            db.rollback()
        else:
            db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    verb = "Would add" if dry_run else "Added"
    if categories_created:
        cverb = "Would create" if dry_run else "Created"
        print(f"{cverb} {len(categories_created)} new vendor categor{'y' if len(categories_created) == 1 else 'ies'}:")
        for c in categories_created:
            print(f"  + {c}")
        print()

    print(f"{verb} {len(added)} vendor(s):")
    for name, cat in added:
        print(f"  + {name}  [{cat or 'no category'}]")
    print(f"\nSkipped {len(skipped)} (already present):")
    for name, cat in skipped:
        print(f"  = {name}  [{cat or 'no category'}]")

    if contact_warnings:
        print(f"\n{len(contact_warnings)} vendor(s) had more than one contact number on the sheet - only the first was kept (the Contact Number field holds one number):")
        for name, original, kept in contact_warnings:
            print(f"  ! {name}: kept \"{kept}\", sheet had \"{original}\"")

    if dry_run:
        print("\nDry run only - no changes were written. Re-run without --dry-run to apply.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", help="Path to the vendor list .xlsx file")
    parser.add_argument("--dry-run", action="store_true", help="Preview changes without writing to the database.")
    args = parser.parse_args()
    main(args.xlsx_path, dry_run=args.dry_run)
